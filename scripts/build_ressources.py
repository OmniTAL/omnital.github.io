#!/usr/bin/env python3
"""Construit data/ressources.json à partir des classeurs Excel de sources/.

- lit la feuille « maîtresse » de chaque classeur (les autres feuilles sont des vues filtrées) ;
- normalise catégories, licences et liens ; dédoublonne par URL ;
- attribue un identifiant stable (hash du lien) ;
- géolocalise uniquement ce qui a une géographie réelle (variété berbère, pays).

Usage : python scripts/build_ressources.py
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from geo import detect_region, geo_fields, region_from_country  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
SOURCES_DIR = ROOT / "sources"
OUTPUT = ROOT / "data" / "ressources.json"

# id source -> (fichier, feuille maîtresse, libellé)
WORKBOOKS = [
    ("berbere",    "nlp_berbere_kabyle_ressources.xlsx",       "Toutes_ressources", "TAL berbère & kabyle"),
    ("ecosysteme", "nlp_depots_datasets_modeles_outils.xlsx",  "Tous_depots",       "Écosystème : dépôts, datasets, modèles, outils"),
    ("general",    "nlp_ressources_filtrees.xlsx",             "Toutes_donnees",    "Veille générale : revues, événements, emplois"),
]

CATEGORY_RULES = [
    (re.compile(r"dataset|corpus", re.I),             "cat_Dataset"),
    (re.compile(r"mod[eè]le", re.I),                  "cat_Model"),
    (re.compile(r"outil|d[eé]p[oô]t", re.I),          "cat_Tool"),
    (re.compile(r"revue", re.I),                      "cat_Journal"),
    (re.compile(r"conf[eé]rence|[eé]v[eé]nement", re.I), "cat_Event"),
    (re.compile(r"emploi|stage|freelance", re.I),     "cat_Job"),
]

# Coquilles relevées dans les classeurs
TYPO_FIXES = {
    "Confé·ırences": "Conférences",
    "Evements": "Événements",
    "plutot": "plutôt",
    "Conferences": "Conférences",
    "journees": "journées",
    "generative": "générative",
    "verifiees": "vérifiées",
    "localises": "localisés",
    "qualité±±": "qualité variable",
}

LICENCE_CLASSES = [
    ("restricted", re.compile(r"^non|licenci|s[eé]lection|restreint", re.I)),
    ("partial",    re.compile(r"partiel|hybride|variable|compte requis", re.I)),
    ("open",       re.compile(r"ouvert|^oui|open|cc0|cc-by|mit|apache|bsd|gpl", re.I)),
]


def clean(value) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", " ", str(value)).strip()
    if not s or s.upper() in {"N/A", "NA", "NONE", "-"}:
        return None
    for bad, good in TYPO_FIXES.items():
        s = s.replace(bad, good)
    return s


def normalize_url(url: str) -> str:
    u = url.strip()
    u = re.sub(r"^http://(?!www\.wikicfp)", "https://", u)  # wikicfp n'a pas de HTTPS fiable
    return u


def url_key(url: str) -> str:
    return re.sub(r"^https?://(www\.)?", "", url.lower()).rstrip("/")


def stable_id(url: str) -> str:
    return "res-" + hashlib.sha1(url_key(url).encode()).hexdigest()[:10]


def category_key(*labels: str | None) -> str:
    blob = " ".join(l for l in labels if l)
    for rx, key in CATEGORY_RULES:
        if rx.search(blob):
            return key
    return "cat_Tool"


def licence_class(licence: str | None) -> str:
    if not licence:
        return "unknown"
    for cls, rx in LICENCE_CLASSES:
        if rx.search(licence):
            return cls
    return "unknown"


def read_sheet(path: Path, sheet: str) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]


def build_item(source_id: str, row: dict) -> dict | None:
    nom, lien = clean(row.get("nom")), clean(row.get("lien"))
    if not nom or not lien:
        return None
    lien = normalize_url(lien)
    categorie = clean(row.get("categorie"))
    rtype = clean(row.get("type"))
    langue = clean(row.get("langue"))
    pays = clean(row.get("pays"))
    licence = clean(row.get("licence")) or clean(row.get("open_access"))
    description = clean(row.get("description")) or rtype or ""

    # Géographie : variété berbère d'abord, puis pays déclaré
    region = None
    if source_id == "berbere":
        region = detect_region(langue, nom, description)
    region = region or region_from_country(pays)

    return {
        "id": stable_id(lien),
        "nom": nom,
        "source": source_id,
        "categorie": categorie or rtype,
        "categorie_key": category_key(categorie, rtype if not categorie else None),
        "type": rtype,
        "langue": langue,
        "taille": clean(row.get("taille")) if clean(row.get("taille")) != "Variable" else None,
        "licence": licence,
        "licence_class": licence_class(licence),
        "lien": lien,
        "description": {"fr": description},
        "pays": pays or ("International" if region in (None, "International") else None),
        **geo_fields(region),
    }


def main() -> int:
    items: dict[str, dict] = {}
    sources_meta = []
    for source_id, filename, sheet, label in WORKBOOKS:
        path = SOURCES_DIR / filename
        if not path.exists():
            print(f"[WARN] classeur absent : {path.name}")
            continue
        count = 0
        for row in read_sheet(path, sheet):
            item = build_item(source_id, row)
            if not item:
                continue
            key = url_key(item["lien"])
            if key in items:  # doublon inter-classeurs : on complète les champs vides
                for k, v in item.items():
                    if items[key].get(k) in (None, "", {}) and v:
                        items[key][k] = v
                continue
            items[key] = item
            count += 1
        sources_meta.append({"id": source_id, "label": label, "file": filename, "count": count})
        print(f"[OK] {filename} → {count} ressources")

    ordered = sorted(items.values(), key=lambda r: (r["source"], r["categorie_key"], r["nom"].lower()))
    by_cat: dict[str, int] = {}
    for r in ordered:
        by_cat[r["categorie_key"]] = by_cat.get(r["categorie_key"], 0) + 1

    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "total": len(ordered),
        "sources": sources_meta,
        "counts": {"by_source": {s["id"]: s["count"] for s in sources_meta}, "by_category": by_cat},
        "items": ordered,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    new_text = json.dumps(payload, ensure_ascii=False, indent=1)
    # Évite un commit quotidien inutile : on ne réécrit que si le contenu a changé
    if OUTPUT.exists():
        old = json.loads(OUTPUT.read_text(encoding="utf-8"))
        old.pop("generated_at", None)
        cmp = dict(payload); cmp.pop("generated_at")
        if old == cmp:
            print(f"[=] {OUTPUT.relative_to(ROOT)} inchangé ({len(ordered)} ressources)")
            return 0
    OUTPUT.write_text(new_text, encoding="utf-8")
    geo = sum(1 for r in ordered if r["lat"] is not None)
    print(f"[OK] {OUTPUT.relative_to(ROOT)} : {len(ordered)} ressources ({geo} géolocalisées)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
